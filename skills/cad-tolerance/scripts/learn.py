#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""learn.py — NAP TAI LIEU vao kho tri thuc cua agent AutoCAD (chay tai may dich).

Tai lieu o dau thi hoc o do: tro script vao file/thu muc tai lieu tren may do,
no trich text (txt/md/csv/pdf/docx/anh-OCR) roi luu vao <repo>/knowledge/.

Dung:
  python learn.py add "D:\\tai_lieu\\giao_trinh.pdf"
  python learn.py add "D:\\tai_lieu\\" --recursive
  python learn.py search "nham be mat"
  python learn.py list
  python learn.py stats
  python learn.py --selftest
"""
import argparse
import glob
import hashlib
import json
import os
import sys
import time
import unicodedata

try:                      # console Windows cp1252 -> khong in duoc dau tieng Viet
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", ".."))
KNOW = os.environ.get("CAD_KNOWLEDGE") or os.path.join(REPO, "knowledge")
INDEX = os.path.join(KNOW, "index.json")

TEXT_EXT = {".txt", ".md", ".csv", ".json", ".log", ".py", ".ini", ".yml", ".yaml"}
IMG_EXT = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}
DOC_EXT = {".pdf", ".docx"}


def nodau(s):
    s = unicodedata.normalize("NFD", str(s).lower())
    return "".join(c for c in s if unicodedata.category(c) != "Mn").replace("đ", "d")


def need(pkg, hint):
    print("  ! Thieu thu vien '%s' -> bo qua dang file nay. Cai bang: %s" % (pkg, hint))
    return None


def extract(path):
    """-> (text, note). note ghi ro da lam gi / thieu gi."""
    ext = os.path.splitext(path)[1].lower()
    if ext in TEXT_EXT:
        with open(path, encoding="utf-8", errors="ignore") as f:
            return f.read(), "doc truc tiep"
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            return None, need("pypdf", "pip install pypdf")
        r = PdfReader(path)
        pages = [(p.extract_text() or "") for p in r.pages]
        txt = "\n\n".join(pages)
        note = "pypdf %d trang" % len(pages)
        if len(txt.strip()) < 50 * len(pages):
            note += " | CO VE LA BAN SCAN -> can OCR (learn.py chua OCR PDF, hay chuyen PDF thanh anh)"
        return txt, note
    if ext == ".docx":
        try:
            import docx
        except ImportError:
            return None, need("python-docx", "pip install python-docx")
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts), "python-docx"
    if ext in IMG_EXT:
        try:
            import easyocr
        except ImportError:
            return None, need("easyocr", "pip install easyocr")
        rd = get_reader()
        lines = rd.readtext(path, detail=0, paragraph=True)
        return "\n".join(lines), "easyocr (vi+en)"
    return None, "dinh dang khong ho tro (%s)" % ext


_READER = None


def get_reader():
    global _READER
    if _READER is None:
        import easyocr
        _READER = easyocr.Reader(["vi", "en"], gpu=False, verbose=False)
    return _READER


def load_index():
    if os.path.exists(INDEX):
        with open(INDEX, encoding="utf-8") as f:
            return json.load(f)
    return {"docs": []}


def save_index(ix):
    os.makedirs(KNOW, exist_ok=True)
    with open(INDEX, "w", encoding="utf-8") as f:
        json.dump(ix, f, ensure_ascii=False, indent=2)


def slug_for(path):
    base = os.path.splitext(os.path.basename(path))[0]
    base = "".join(c if c.isalnum() or c in " -_" else "_" for c in base).strip().replace(" ", "_")
    h = hashlib.md5(os.path.abspath(path).encode()).hexdigest()[:6]
    return "%s_%s" % (base[:60], h)


def add_one(path, ix):
    if not os.path.isfile(path):
        print("  - bo qua (khong phai file): %s" % path)
        return False
    text, note = extract(path)
    if not text or not text.strip():
        print("  x %s -> %s" % (os.path.basename(path), note))
        return False
    slug = slug_for(path)
    out = os.path.join(KNOW, slug + ".md")
    os.makedirs(KNOW, exist_ok=True)
    with open(out, "w", encoding="utf-8") as f:
        f.write("<!-- nguon: %s -->\n<!-- nap luc: %s -->\n<!-- cach trich: %s -->\n\n" % (
            path, time.strftime("%Y-%m-%d %H:%M"), note))
        f.write(text)
    ix["docs"] = [d for d in ix["docs"] if d["file"] != slug + ".md"]
    ix["docs"].append({
        "file": slug + ".md", "nguon": os.path.abspath(path),
        "chars": len(text), "cach_trich": note, "nap_luc": time.strftime("%Y-%m-%d %H:%M"),
    })
    print("  + %s -> knowledge/%s.md  (%d ky tu | %s)" % (os.path.basename(path), slug, len(text), note))
    return True


def cmd_add(a):
    targets = []
    if a.recursive and os.path.isdir(a.path):
        for ext in TEXT_EXT | IMG_EXT | DOC_EXT:
            targets += glob.glob(os.path.join(a.path, "**", "*" + ext), recursive=True)
    else:
        targets = [a.path]
    if not targets:
        print("Khong tim thay file nao.")
        return 1
    ix = load_index()
    ok = 0
    for t in targets:
        if add_one(t, ix):
            ok += 1
    save_index(ix)
    print("Da hoc %d/%d tai lieu. Kho: %s" % (ok, len(targets), os.path.join(KNOW, "index.json")))
    return 0 if ok else 1


def cmd_search(a):
    if not os.path.isdir(KNOW):
        print("Kho tri thuc rong: %s" % KNOW)
        return 1
    key = nodau(a.query)
    hits = 0
    for fn in sorted(glob.glob(os.path.join(KNOW, "*.md"))):
        text = open(fn, encoding="utf-8", errors="ignore").read()
        low = nodau(text)
        pos = low.find(key)
        while pos != -1:
            hits += 1
            s = max(0, pos - 120)
            print("--- %s @ %d ---" % (os.path.basename(fn), pos))
            print(text[s:pos + 200].replace("\n", " ").strip())
            if hits >= a.limit:
                print("\n(dung o %d ket qua, tang --limit de xem them)" % hits)
                return 0
            pos = low.find(key, pos + 1)
    if not hits:
        print("Khong tim thay '%s' trong kho." % a.query)
        return 1
    return 0


def cmd_list(a):
    ix = load_index()
    if not ix["docs"]:
        print("Kho rong — chua hoc tai lieu nao. Dung: python learn.py add \"<duong dan>\"")
        return 0
    for d in ix["docs"]:
        print("%-42s %8d ky tu  %s" % (d["file"], d["chars"], d["nguon"]))
    return 0


def cmd_stats(a):
    ix = load_index()
    print("Kho        : %s" % KNOW)
    print("So tai lieu : %d" % len(ix["docs"]))
    print("Tong ky tu  : %d" % sum(d["chars"] for d in ix["docs"]))
    return 0


def selftest():
    import tempfile
    global KNOW, INDEX
    old = (KNOW, INDEX)
    tmp = tempfile.mkdtemp(prefix="cadknow_")
    KNOW = tmp
    INDEX = os.path.join(tmp, "index.json")
    try:
        src = os.path.join(tmp, "src")
        os.makedirs(src)
        p = os.path.join(src, "thu.txt")
        with open(p, "w", encoding="utf-8") as f:
            f.write("Ra la do nham be mat. Cap chinh xac IT7 cho lo lap o lan.")
        ix = load_index()
        assert add_one(p, ix) is True, "phai nap duoc file txt"
        save_index(ix)
        assert os.path.exists(os.path.join(KNOW, slug_for(p) + ".md")), "phai sinh file .md"
        assert load_index()["docs"][0]["chars"] > 10, "index phai ghi so ky tu"
    finally:
        KNOW, INDEX = old
        import shutil
        shutil.rmtree(tmp, ignore_errors=True)
    print("SELFTEST PASS")
    return 0


def main():
    p = argparse.ArgumentParser(description="Nap tai lieu vao kho tri thuc agent AutoCAD")
    sub = p.add_subparsers(dest="cmd")
    a1 = sub.add_parser("add", help="nap file/thu muc tai lieu")
    a1.add_argument("path")
    a1.add_argument("--recursive", action="store_true")
    a1.set_defaults(fn=cmd_add)
    a2 = sub.add_parser("search", help="tim trong kho da hoc")
    a2.add_argument("query")
    a2.add_argument("--limit", type=int, default=5)
    a2.set_defaults(fn=cmd_search)
    a3 = sub.add_parser("list", help="liet ke tai lieu da hoc")
    a3.set_defaults(fn=cmd_list)
    a4 = sub.add_parser("stats", help="thong ke kho")
    a4.set_defaults(fn=cmd_stats)
    p.add_argument("--selftest", action="store_true")
    a = p.parse_args()
    if a.selftest:
        return selftest()
    if not getattr(a, "fn", None):
        p.print_help()
        return 1
    return a.fn(a)


if __name__ == "__main__":
    sys.exit(main())
